#!/usr/bin/env python3
"""Build the GBE initial-submission DOCX package from audited Markdown sources."""

from __future__ import annotations

import re
import shutil
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
MANUSCRIPT = ROOT / "manuscript"
FIGURES = ROOT / "figures" / "manuscript"
OUT = ROOT / "submission" / "GBE_initial_submission_2026-10-06"
DATA_OUTPUT = ROOT / "outputs" / "gbe_supplementary_data_2026-10-06"

MAIN_SOURCE = MANUSCRIPT / "GBE_SUBMISSION_DRAFT.md"
SUPP_SOURCE = MANUSCRIPT / "GBE_SUPPLEMENTARY_TEXT.md"
COVER_SOURCE = MANUSCRIPT / "GBE_COVER_LETTER_DRAFT.md"

MAIN_FIGURES = {
    "1": FIGURES / "Figure_1_regional_constraint.png",
    "2": FIGURES / "Figure_2_transcript_architecture.png",
    "3": FIGURES / "Figure_3_developmental_isoform_function.png",
    "4": FIGURES / "Figure_4_neuronal_mCA.png",
}
SUPP_FIGURES = {
    "S1": FIGURES / "Figure_S1_chromatin_context.png",
    "S2": FIGURES / "Figure_S2_L188H_preparation_sensitivity.png",
}


def set_run_font(run, name: str = "Times New Roman", size: float = 12) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    rfonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), name)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    set_run_font(run, size=10)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, text, end])


def configure_document(doc: Document, *, line_numbers: bool, double_spaced: bool) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)

    if line_numbers:
        sect_pr = section._sectPr
        line_num = sect_pr.find(qn("w:lnNumType"))
        if line_num is None:
            line_num = OxmlElement("w:lnNumType")
            sect_pr.append(line_num)
        line_num.set(qn("w:countBy"), "1")
        line_num.set(qn("w:start"), "1")
        line_num.set(qn("w:restart"), "continuous")
        line_num.set(qn("w:distance"), "360")

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.paragraph_format.line_spacing = 2 if double_spaced else 1.15
    normal.paragraph_format.space_after = Pt(0 if double_spaced else 6)

    for style_name, size in (("Title", 14), ("Heading 1", 13), ("Heading 2", 12)):
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
        p_pr = style._element.get_or_add_pPr()
        p_bdr = p_pr.find(qn("w:pBdr"))
        if p_bdr is not None:
            p_pr.remove(p_bdr)

    add_page_number(section.footer.paragraphs[0])


INLINE = re.compile(r"(\*\*.*?\*\*|\*.*?\*|`.*?`)")


def add_inline(paragraph, text: str) -> None:
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    for part in INLINE.split(text):
        if not part:
            continue
        bold = part.startswith("**") and part.endswith("**")
        italic = (part.startswith("*") and part.endswith("*")) and not bold
        code = part.startswith("`") and part.endswith("`")
        content = part[2:-2] if bold else part[1:-1] if (italic or code) else part
        run = paragraph.add_run(content)
        set_run_font(run)
        run.bold = bold
        run.italic = italic
        if code:
            set_run_font(run, "Courier New", 10.5)


def set_picture_alt_text(run, title: str, description: str) -> None:
    doc_pr = run._element.xpath(".//wp:docPr")
    if doc_pr:
        doc_pr[0].set("title", title)
        doc_pr[0].set("descr", re.sub(r"\*+", "", description))


def figure_alt_text(markdown: str, number: str) -> str:
    pattern = re.compile(
        rf"\*\*Fig\. {re.escape(number)}\..*?\*\*.*?\n\n\*\*Alt text:\*\* (.*?)(?:\n\n|$)",
        re.S,
    )
    match = pattern.search(markdown)
    return " ".join(match.group(1).split()) if match else f"Figure {number}"


def add_figure(doc: Document, number: str, image_path: Path, alt: str, width: float = 6.4) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_together = True
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Inches(width))
    set_picture_alt_text(run, f"Figure {number}", alt)
    label = doc.add_paragraph()
    label.alignment = WD_ALIGN_PARAGRAPH.CENTER
    label.paragraph_format.keep_with_next = False
    label.paragraph_format.space_after = Pt(6)
    r = label.add_run(f"Figure {number}")
    set_run_font(r, size=11)
    r.bold = True


def add_markdown(doc: Document, markdown: str, *, main_figures: bool = False) -> None:
    blocks = re.split(r"\n\s*\n", markdown.strip())
    for block in blocks:
        block = block.strip()
        if not block:
            continue
        marker = re.fullmatch(r"\*\*\[FIGURE ([1-4]) NEAR HERE\]\*\*", block)
        if marker and main_figures:
            number = marker.group(1)
            add_figure(doc, number, MAIN_FIGURES[number], figure_alt_text(markdown, number))
            continue
        if block.startswith("# "):
            paragraph = doc.add_paragraph(style="Title")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline(paragraph, block[2:].strip())
        elif block.startswith("## "):
            add_inline(doc.add_paragraph(style="Heading 1"), block[3:].strip())
        elif block.startswith("### "):
            add_inline(doc.add_paragraph(style="Heading 2"), block[4:].strip())
        elif all(line.startswith("- ") for line in block.splitlines()):
            for line in block.splitlines():
                add_inline(doc.add_paragraph(style="List Bullet"), line[2:])
        else:
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.widow_control = True
            add_inline(paragraph, " ".join(line.strip() for line in block.splitlines()))


def build_main() -> Path:
    markdown = MAIN_SOURCE.read_text()
    doc = Document()
    configure_document(doc, line_numbers=True, double_spaced=True)
    add_markdown(doc, markdown, main_figures=True)
    path = OUT / "GBE_Main_Manuscript.docx"
    doc.save(path)
    return path


def build_cover() -> Path:
    markdown = COVER_SOURCE.read_text().replace("[Date]", "October 6, 2026")
    markdown = markdown.replace(
        "The manuscript is\nnot under consideration elsewhere, and all authors must approve the submitted\nversion. [Confirm or revise these statements before submission.]",
        "Upon submission, the corresponding author will confirm that the manuscript is\nnot under consideration elsewhere and that all authors approve the submitted\nversion.",
    )
    markdown = markdown.replace(
        "# Cover letter: Genome Biology and Evolution",
        "# Cover Letter for Genome Biology and Evolution",
    )
    doc = Document()
    configure_document(doc, line_numbers=False, double_spaced=False)
    add_markdown(doc, markdown)
    path = OUT / "GBE_Cover_Letter.docx"
    doc.save(path)
    return path


def supplementary_legend(markdown: str, number: str) -> tuple[str, str]:
    pattern = re.compile(
        rf"(\*\*Fig\. {re.escape(number)}\..*?\*\*.*?)(?:\n\n\*\*Alt text:\*\* (.*?))(?:\n\n|$)",
        re.S,
    )
    match = pattern.search(markdown)
    if not match:
        raise ValueError(f"Missing supplementary legend for {number}")
    return " ".join(match.group(1).split()), " ".join(match.group(2).split())


def build_supplement() -> Path:
    markdown = SUPP_SOURCE.read_text()
    doc = Document()
    configure_document(doc, line_numbers=False, double_spaced=True)
    title = markdown.splitlines()[0].removeprefix("# ").replace(
        "Supplementary information for:", "Supplementary Information for"
    )
    title_p = doc.add_paragraph(style="Title")
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline(title_p, title)
    author_p = doc.add_paragraph()
    author_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline(author_p, "Jie Lu")

    data_tables_match = re.search(
        r"(## Supplementary data tables.*?)(?=\n## Exploratory structural analysis)",
        markdown,
        re.S,
    )
    if not data_tables_match:
        raise ValueError("Missing supplementary data-table list")
    add_markdown(doc, data_tables_match.group(1))

    doc.add_page_break()
    legend, alt = supplementary_legend(markdown, "S1")
    add_figure(doc, "S1", SUPP_FIGURES["S1"], alt, width=6.4)
    paragraph = doc.add_paragraph()
    add_inline(paragraph, legend)
    alt_p = doc.add_paragraph()
    add_inline(alt_p, f"**Alt text:** {alt}")

    analysis_match = re.search(
        r"(## Exploratory structural analysis of L188H.*?)(?=\n## Supplementary figure legends)",
        markdown,
        re.S,
    )
    if not analysis_match:
        raise ValueError("Missing supplementary L188H analysis")
    doc.add_page_break()
    analysis_start = len(doc.paragraphs)
    add_markdown(doc, analysis_match.group(1))
    for paragraph in doc.paragraphs[analysis_start:]:
        paragraph.paragraph_format.line_spacing = 1.5

    doc.add_page_break()
    legend, alt = supplementary_legend(markdown, "S2")
    add_figure(doc, "S2", SUPP_FIGURES["S2"], alt, width=6.4)
    paragraph = doc.add_paragraph()
    add_inline(paragraph, legend)
    alt_p = doc.add_paragraph()
    add_inline(alt_p, f"**Alt text:** {alt}")
    path = OUT / "GBE_Supplementary_Information.docx"
    doc.save(path)
    return path


def copy_figure_sources() -> None:
    destination = OUT / "figure_source_files_for_revision"
    destination.mkdir(parents=True, exist_ok=True)
    for stem in (
        "Figure_1_regional_constraint",
        "Figure_2_transcript_architecture",
        "Figure_3_developmental_isoform_function",
        "Figure_4_neuronal_mCA",
        "Figure_S1_chromatin_context",
        "Figure_S2_L188H_preparation_sensitivity",
    ):
        for suffix in (".pdf", ".svg", ".png"):
            shutil.copy2(FIGURES / f"{stem}{suffix}", destination / f"{stem}{suffix}")


def copy_supplementary_data() -> None:
    for name in (
        "GBE_Supplementary_Data_Tables.xlsx",
        "GBE_Supplementary_Machine_Readable_Data.zip",
    ):
        source = DATA_OUTPUT / name
        if not source.exists():
            raise FileNotFoundError(source)
        shutil.copy2(source, OUT / name)


def copy_markdown_sources() -> None:
    for source in (MAIN_SOURCE, SUPP_SOURCE, COVER_SOURCE):
        shutil.copy2(source, OUT / source.name)


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    paths = [build_main(), build_cover(), build_supplement()]
    copy_figure_sources()
    copy_supplementary_data()
    copy_markdown_sources()
    for path in paths:
        print(path)


if __name__ == "__main__":
    main()
